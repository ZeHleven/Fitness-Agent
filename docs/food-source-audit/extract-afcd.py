"""Read-only extraction of the official AFCD source workbooks; never rewrites XLSX."""
import hashlib
import json
from pathlib import Path
import openpyxl

out = Path(__file__).resolve().parent
base = out.parents[2] / 'food-library-planning-20260913' / 'sources'
files = [('afcd-food-details.xlsx', 'Food details'),
         ('afcd-nutrient-profiles.xlsx', 'All solids & liquids per 100 g')]
evidence = {'sourceId': 'F009805', 'retrieved': '2026-09-13', 'sources': []}
for filename, sheetname in files:
    path = base / filename
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb[sheetname]
    headers = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    hits = [(i, r) for i, r in enumerate(sheet.iter_rows(values_only=True), 1) if r[0] == 'F009805']
    assert len(hits) == 1
    index, row = hits[0]
    columns = range(11) if filename.endswith('details.xlsx') else [0, 2, 3, 4, 5, 7, 9, 11, 38, 39]
    evidence['sources'].append({'file': filename, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'sheet': sheetname, 'row': index, 'fields': [
            {'cell': f'{openpyxl.utils.get_column_letter(c+1)}{index}', 'label': headers[c], 'raw': row[c]} for c in columns]})
    wb.close()
assert evidence['sources'][0]['fields'][4]['raw'] == 'Uncooked duck breast meat with tendons, skin and fat removed'
assert [x['raw'] for x in evidence['sources'][1]['fields'][3:]] == [383, 383, 21.2, 0.6, 0, 0, 0]
(out / 'afcd-duck-evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
print(json.dumps(evidence, ensure_ascii=False, indent=2))
